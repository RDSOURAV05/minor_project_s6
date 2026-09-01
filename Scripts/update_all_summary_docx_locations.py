import os
import shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from swap_base_paper import papers_data

WORKSPACE = r"c:\Users\PRO\OneDrive\Documents\GitHub\minor project"

def set_cell_bg(cell, hex_col):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_col)
    tcPr.append(shd)

def generate_master_docx():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8); s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8); s.right_margin = Inches(0.8)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_title.add_run("IEEE Literature Survey & Summary Reference Document")
    r.font.name = 'Calibri'; r.font.size = Pt(22); r.font.bold = True; r.font.color.rgb = RGBColor(0, 51, 102)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("AI-Based Audio Watermark Detection for Copyright Protection and Deepfake Authentication\n12 Peer-Reviewed IEEE Journal Papers (2020–2026) | Base Paper: DeepMark Benchmark")
    r_sub.font.name = 'Calibri'; r_sub.font.size = Pt(11); r_sub.font.italic = True; r_sub.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    h1 = doc.add_paragraph()
    r_h1 = h1.add_run("1. Executive Summary Table of Reviewed Papers")
    r_h1.font.name = 'Calibri'; r_h1.font.size = Pt(14); r_h1.font.bold = True; r_h1.font.color.rgb = RGBColor(0, 51, 102)

    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    headers = ["#", "Paper Title & IEEE Journal", "Year", "Core Methodology", "Key Metrics", "DOI Link"]
    for idx, text in enumerate(headers):
        hdr_cells[idx].text = text
        set_cell_bg(hdr_cells[idx], "003366")
        p = hdr_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = 'Calibri'; r.font.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(255, 255, 255)

    for p_item in papers_data:
        row_cells = table.add_row().cells
        b_tag = " (BASE PAPER)" if p_item['id'] == 1 else ""
        row_cells[0].text = str(p_item['id'])
        row_cells[1].text = f"{p_item['title']}{b_tag}\n({p_item['journal']})"
        row_cells[2].text = str(p_item['year'])
        row_cells[3].text = p_item['methodology']
        row_cells[4].text = p_item['metrics']
        row_cells[5].text = p_item['doi']
        for cell in row_cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
                for r in p.runs:
                    r.font.name = 'Calibri'; r.font.size = Pt(8.5)

    doc.add_page_break()

    h2 = doc.add_paragraph()
    r_h2 = h2.add_run("2. Detailed Paper Summaries & Reference Mapping")
    r_h2.font.name = 'Calibri'; r_h2.font.size = Pt(14); r_h2.font.bold = True; r_h2.font.color.rgb = RGBColor(0, 51, 102)

    for p_item in papers_data:
        b_tag = " (BASE PAPER)" if p_item['id'] == 1 else ""
        p_hdr = doc.add_paragraph()
        r_phdr = p_hdr.add_run(f"Paper {p_item['id']}{b_tag}: {p_item['title']}")
        r_phdr.font.name = 'Calibri'; r_phdr.font.size = Pt(13); r_phdr.font.bold = True; r_phdr.font.color.rgb = RGBColor(0, 85, 128)

        p_meta = doc.add_paragraph()
        p_meta.add_run(f"Authors: {p_item['authors']}\n").bold = True
        p_meta.add_run(f"Journal: {p_item['journal']} ({p_item['year']}) | Vol: {p_item['volume']}, Pages: {p_item['pages']}\n")
        p_meta.add_run(f"DOI: https://doi.org/{p_item['doi']}\n")
        p_meta.add_run(f"Local PDF File: {p_item['filename']}")
        p_meta.paragraph_format.space_after = Pt(6)

        p_ab = doc.add_paragraph()
        p_ab.add_run("Abstract Summary: ").bold = True
        p_ab.add_run(p_item['abstract'])
        p_ab.paragraph_format.space_after = Pt(4)

        p_m = doc.add_paragraph()
        p_m.add_run("Methodology & Features: ").bold = True
        p_m.add_run(p_item['methodology'])
        p_m.paragraph_format.space_after = Pt(4)

        p_d = doc.add_paragraph()
        p_d.add_run("Datasets Evaluated: ").bold = True
        p_d.add_run(p_item['datasets'])
        p_d.paragraph_format.space_after = Pt(4)

        p_res = doc.add_paragraph()
        p_res.add_run("Reported Results & Accuracy: ").bold = True
        p_res.add_run(p_item['metrics'])
        p_res.paragraph_format.space_after = Pt(4)

        p_xy = doc.add_paragraph()
        p_xy.add_run("Graph Parameters (X & Y Axes): ").bold = True
        p_xy.add_run(p_item['graph_x_y'])
        p_xy.paragraph_format.space_after = Pt(4)

        p_pro = doc.add_paragraph()
        p_pro.add_run("Pros: ").bold = True
        p_pro.add_run(p_item['pros'] + " | ")
        p_pro.add_run("Cons: ").bold = True
        p_pro.add_run(p_item['cons'])
        p_pro.paragraph_format.space_after = Pt(4)

        p_cite = doc.add_paragraph()
        p_cite.add_run("IEEE Citation: ").bold = True
        p_cite.add_run(f"{p_item['authors']}, \"{p_item['title']},\" {p_item['journal']}, vol. {p_item['volume']}, pp. {p_item['pages']}, {p_item['year']}, doi: {p_item['doi']}.")
        p_cite.paragraph_format.space_after = Pt(12)

        doc.add_paragraph("-" * 80).paragraph_format.space_after = Pt(12)

    # Save to all 3 paths
    paths = [
        os.path.join(WORKSPACE, "IEEE_Papers_Summary.docx"),
        os.path.join(WORKSPACE, "Documentation", "IEEE_Papers_Summary.docx"),
        os.path.join(WORKSPACE, "IEEE_Research_Papers", "Summaries", "IEEE_Papers_Summary.docx")
    ]
    for p_out in paths:
        os.makedirs(os.path.dirname(p_out), exist_ok=True)
        doc.save(p_out)
        print(f"Saved summary docx: {p_out}")

if __name__ == "__main__":
    generate_master_docx()
